import json
import os
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))


def sample_deck(n_slides=9, n_refs=11):
    return {
        "title": "Retrieval-Augmented Generation: State of Evaluation",
        "subtitle": "What 50 papers say about measuring RAG quality",
        "topic": "retrieval-augmented generation evaluation",
        "generated_on": "2026-09-19",
        "stats_line": "Synthesized from 50 papers (31 full-text, 19 abstract-only) via Semantic Scholar",
        "slides": [
            {
                "title": f"Finding {i + 1}",
                "bullets": [{"text": f"Bullet {j + 1} with a specific finding about RAG.", "refs": [1 + (i + j) % n_refs]} for j in range(4)],
                "notes": f"Speaker notes for slide {i + 1}.",
            }
            for i in range(n_slides)
        ],
        "references": [{"n": k + 1, "text": f"Author {k + 1} et al. (2024). Paper {k + 1}. Venue."} for k in range(n_refs)],
    }


class RenderDeckTest(unittest.TestCase):
    def render(self, deck):
        out = os.path.join(tempfile.mkdtemp(), "deck.pptx")
        subprocess.run(
            [sys.executable, os.path.join(HERE, "render_deck.py"), out],
            input=json.dumps(deck).encode("utf-8"),
            check=True,
        )
        return Presentation(out)

    def test_structure_notes_and_citations(self):
        deck = sample_deck()
        prs = self.render(deck)
        # title + 9 content + 2 reference slides (11 refs at 8 per slide)
        self.assertEqual(len(prs.slides), 1 + 9 + 2)
        for slide in prs.slides:
            self.assertTrue(slide.notes_slide.notes_text_frame.text.strip(), "every slide needs speaker notes")
        content = list(prs.slides)[1:10]
        for slide in content:
            text = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
            self.assertEqual(text.count("▸"), 4)
            self.assertRegex(text, r"\[\d+\]")
        ref_text = " ".join(sh.text_frame.text for s in list(prs.slides)[10:] for sh in s.shapes if sh.has_text_frame)
        for k in range(1, 12):
            self.assertIn(f"[{k}]", ref_text)

    def test_unicode_titles(self):
        deck = sample_deck(n_slides=8, n_refs=3)
        deck["title"] = "Évaluation — RAG «multilingue» 検索"
        prs = self.render(deck)
        self.assertEqual(len(prs.slides), 1 + 8 + 1)


if __name__ == "__main__":
    unittest.main()
