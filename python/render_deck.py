"""Render a research deck to PPTX with python-pptx.

Usage: python render_deck.py OUTPUT.pptx < deck.json

The input JSON matches RenderDeck in src/lib/deckSchema.ts. Branding (colors, fonts,
footer, optional logo) comes from brand.json next to this script, or BRAND_CONFIG.
"""

import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.7)
REFS_PER_SLIDE = 8


def load_brand():
    path = os.environ.get("BRAND_CONFIG") or os.path.join(HERE, "brand.json")
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


class DeckBuilder:
    def __init__(self, brand):
        self.brand = brand
        self.c = brand["colors"]
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.page = 0

    # ---- primitives -------------------------------------------------------
    def _rect(self, slide, left, top, width, height, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def _text(self, slide, left, top, width, height, text, size, color, font=None, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Emu(0)
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        self._style(run, size, color, font or self.brand["fonts"]["body"], bold)
        return box

    @staticmethod
    def _style(run, size, color, font, bold=False):
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
        run.font.name = font
        run.font.bold = bold

    def _notes(self, slide, text):
        slide.notes_slide.notes_text_frame.text = text

    def _logo(self, slide):
        logo = self.brand.get("logo_path")
        if logo:
            path = logo if os.path.isabs(logo) else os.path.join(HERE, logo)
            if os.path.exists(path):
                slide.shapes.add_picture(path, SLIDE_W - MARGIN - Inches(1.2), Inches(0.35), height=Inches(0.45))

    def _chrome(self, slide, title):
        """Header bar, title, and footer shared by content slides."""
        self.page += 1
        self._rect(slide, 0, 0, SLIDE_W, Inches(0.12), self.c["primary"])
        self._rect(slide, 0, Inches(0.12), Inches(2.2), Inches(0.05), self.c["accent"])
        self._text(slide, MARGIN, Inches(0.45), SLIDE_W - 2 * MARGIN - Inches(1.4), Inches(0.9), title, 28, self.c["dark"],
                   font=self.brand["fonts"]["heading"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
        self._rect(slide, MARGIN, Inches(1.4), Inches(0.9), Inches(0.04), self.c["primary_light"])
        self._text(slide, MARGIN, SLIDE_H - Inches(0.5), Inches(8), Inches(0.3), self.brand["footer"], 10, self.c["muted"])
        self._text(slide, SLIDE_W - MARGIN - Inches(1), SLIDE_H - Inches(0.5), Inches(1), Inches(0.3), str(self.page), 10,
                   self.c["muted"], align=PP_ALIGN.RIGHT)
        self._logo(slide)

    # ---- slides -----------------------------------------------------------
    def title_slide(self, deck):
        slide = self.prs.slides.add_slide(self.blank)
        self._rect(slide, 0, 0, SLIDE_W, SLIDE_H, self.c["dark"])
        self._rect(slide, MARGIN, Inches(2.3), Inches(1.2), Inches(0.08), self.c["accent"])
        self._text(slide, MARGIN, Inches(1.5), Inches(11), Inches(0.5), self.brand["name"].upper(), 14, self.c["primary_light"], bold=True)
        self._text(slide, MARGIN, Inches(2.6), Inches(11.5), Inches(1.9), deck["title"], 40, "FFFFFF",
                   font=self.brand["fonts"]["heading"], bold=True)
        self._text(slide, MARGIN, Inches(4.2), Inches(11.5), Inches(1.0), deck["subtitle"], 20, self.c["on_dark"])
        self._text(slide, MARGIN, Inches(6.3), Inches(11.5), Inches(0.4), f'{deck["stats_line"]} · {deck["generated_on"]}', 12,
                   self.c["primary_light"])
        self._logo(slide)
        self._notes(slide, f'Topic: {deck["topic"]}\n{deck["stats_line"]}.\nEvery finding in this deck cites its source paper; '
                           f'numbered markers map to the references at the end.')

    def content_slide(self, spec):
        slide = self.prs.slides.add_slide(self.blank)
        self._chrome(slide, spec["title"])
        bullets = spec["bullets"]
        total_chars = sum(len(b["text"]) for b in bullets)
        size = 20 if total_chars < 450 else 18 if total_chars < 650 else 16

        box = slide.shapes.add_textbox(MARGIN, Inches(1.75), SLIDE_W - 2 * MARGIN, Inches(4.9))
        tf = box.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(14)
            marker = p.add_run()
            marker.text = "▸  "
            self._style(marker, size, self.c["primary"], self.brand["fonts"]["body"], bold=True)
            body = p.add_run()
            body.text = bullet["text"]
            self._style(body, size, self.c["text"], self.brand["fonts"]["body"])
            if bullet["refs"]:
                cite = p.add_run()
                cite.text = " [" + ", ".join(str(n) for n in bullet["refs"]) + "]"
                self._style(cite, size - 4, self.c["accent"], self.brand["fonts"]["body"], bold=True)
        self._notes(slide, spec["notes"])

    def references_slides(self, references):
        chunks = [references[i:i + REFS_PER_SLIDE] for i in range(0, len(references), REFS_PER_SLIDE)] or [[]]
        for idx, group in enumerate(chunks):
            slide = self.prs.slides.add_slide(self.blank)
            suffix = f" ({idx + 1}/{len(chunks)})" if len(chunks) > 1 else ""
            self._chrome(slide, f"References{suffix}")
            box = slide.shapes.add_textbox(MARGIN, Inches(1.7), SLIDE_W - 2 * MARGIN, Inches(5.1))
            tf = box.text_frame
            tf.word_wrap = True
            for i, ref in enumerate(group):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(6)
                num = p.add_run()
                num.text = f'[{ref["n"]}]  '
                self._style(num, 12, self.c["primary"], self.brand["fonts"]["body"], bold=True)
                body = p.add_run()
                body.text = ref["text"]
                self._style(body, 12, self.c["text"], self.brand["fonts"]["body"])
            self._notes(slide, "Full references for every numbered citation marker in this deck.")

    def build(self, deck, out_path):
        self.title_slide(deck)
        for spec in deck["slides"]:
            self.content_slide(spec)
        self.references_slides(deck["references"])
        self.prs.save(out_path)


def main():
    if len(sys.argv) != 2:
        print("usage: render_deck.py OUTPUT.pptx < deck.json", file=sys.stderr)
        sys.exit(2)
    deck = json.loads(sys.stdin.buffer.read().decode("utf-8"))
    DeckBuilder(load_brand()).build(deck, sys.argv[1])


if __name__ == "__main__":
    main()
